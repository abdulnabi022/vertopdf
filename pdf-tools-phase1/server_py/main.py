from fastapi import FastAPI, File, UploadFile, Form, BackgroundTasks
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
import shutil
import os
from pathlib import Path
import subprocess
import time
from docx import Document as DocxDocument
from docx.shared import Pt
import PyPDF2
import openpyxl
from pdf2image import convert_from_path
from PIL import Image
import zipfile

app = FastAPI(title="RarePDFtool API", description="All-in-one PDF & image toolkit backend for RarePDFtool.")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

@app.get("/")
def root():
    return {"message": "RarePDFtool server is running and CORS is enabled."}

@app.post("/api/compress")
async def compress_pdf(file: UploadFile = File(...), background_tasks: BackgroundTasks = None):
    input_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
    output_path = UPLOAD_DIR / f"compressed_{int(time.time())}.pdf"
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    gs_args = [
        "gs",
        "-sDEVICE=pdfwrite",
        "-dCompatibilityLevel=1.4",
        "-dPDFSETTINGS=/ebook",
        "-dNOPAUSE",
        "-dQUIET",
        "-dBATCH",
        f"-sOutputFile={output_path}",
        str(input_path)
    ]
    try:
        subprocess.run(gs_args, check=True)
        def cleanup():
            if os.path.exists(input_path): os.remove(input_path)
            if os.path.exists(output_path): os.remove(output_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(output_path, filename="compressed.pdf", media_type="application/pdf")
    except Exception as e:
        if os.path.exists(input_path): os.remove(input_path)
        if os.path.exists(output_path): os.remove(output_path)
        return JSONResponse(status_code=500, content={"error": "Compression failed", "details": str(e)})

@app.post("/api/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...), background_tasks: BackgroundTasks = None):
    input_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
    output_path = UPLOAD_DIR / f"converted_{int(time.time())}.docx"
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    try:
        reader = PyPDF2.PdfReader(str(input_path))
        text = "\n".join(page.extract_text() or '' for page in reader.pages)
        doc = DocxDocument()
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(output_path)
        def cleanup():
            if os.path.exists(input_path): os.remove(input_path)
            if os.path.exists(output_path): os.remove(output_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(output_path, filename="converted.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as e:
        if os.path.exists(input_path): os.remove(input_path)
        if os.path.exists(output_path): os.remove(output_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})

@app.post("/api/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...), background_tasks: BackgroundTasks = None):
    input_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
    output_path = UPLOAD_DIR / f"converted_{int(time.time())}.xlsx"
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    try:
        reader = PyPDF2.PdfReader(str(input_path))
        text = "\n".join(page.extract_text() or '' for page in reader.pages)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PDF Data"
        for line in text.split("\n"):
            ws.append([line])
        wb.save(output_path)
        def cleanup():
            if os.path.exists(input_path): os.remove(input_path)
            if os.path.exists(output_path): os.remove(output_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(output_path, filename="converted.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        if os.path.exists(input_path): os.remove(input_path)
        if os.path.exists(output_path): os.remove(output_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})

@app.post("/api/pdf-to-jpg")
async def pdf_to_jpg(file: UploadFile = File(...), background_tasks: BackgroundTasks = None):
    input_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
    output_dir = UPLOAD_DIR / f"pdf2jpg_{int(time.time())}"
    output_dir.mkdir(exist_ok=True)
    zip_path = UPLOAD_DIR / f"pdf2jpg_{int(time.time())}.zip"
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    try:
        images = convert_from_path(str(input_path), dpi=200, fmt='jpeg')
        jpg_files = []
        for idx, img in enumerate(images, 1):
            jpg_path = output_dir / f"page_{idx}.jpg"
            img.save(jpg_path, "JPEG")
            jpg_files.append(jpg_path)
        with zipfile.ZipFile(zip_path, 'w') as zipf:
            for jpg in jpg_files:
                zipf.write(jpg, arcname=jpg.name)
        def cleanup():
            if os.path.exists(input_path): os.remove(input_path)
            for jpg in jpg_files:
                if os.path.exists(jpg): os.remove(jpg)
            if output_dir.exists():
                try: output_dir.rmdir()
                except: pass
            if os.path.exists(zip_path): os.remove(zip_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(zip_path, filename="pdf-to-jpg.zip", media_type="application/zip")
    except Exception as e:
        if os.path.exists(input_path): os.remove(input_path)
        for jpg in locals().get('jpg_files', []):
            if os.path.exists(jpg): os.remove(jpg)
        if output_dir.exists():
            try: output_dir.rmdir()
            except: pass
        if os.path.exists(zip_path): os.remove(zip_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})

@app.post("/api/pdf-to-png")
async def pdf_to_png(file: UploadFile = File(...), background_tasks: BackgroundTasks = None):
    input_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
    output_dir = UPLOAD_DIR / f"pdf2png_{int(time.time())}"
    output_dir.mkdir(exist_ok=True)
    zip_path = UPLOAD_DIR / f"pdf2png_{int(time.time())}.zip"
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    try:
        images = convert_from_path(str(input_path), dpi=200, fmt='png')
        png_files = []
        for idx, img in enumerate(images, 1):
            png_path = output_dir / f"page_{idx}.png"
            img.save(png_path, "PNG")
            png_files.append(png_path)
        with zipfile.ZipFile(zip_path, 'w') as zipf:
            for png in png_files:
                zipf.write(png, arcname=png.name)
        def cleanup():
            if os.path.exists(input_path): os.remove(input_path)
            for png in png_files:
                if os.path.exists(png): os.remove(png)
            if output_dir.exists():
                try: output_dir.rmdir()
                except: pass
            if os.path.exists(zip_path): os.remove(zip_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(zip_path, filename="pdf-to-png.zip", media_type="application/zip")
    except Exception as e:
        if os.path.exists(input_path): os.remove(input_path)
        for png in locals().get('png_files', []):
            if os.path.exists(png): os.remove(png)
        if output_dir.exists():
            try: output_dir.rmdir()
            except: pass
        if os.path.exists(zip_path): os.remove(zip_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})

@app.post("/api/jpg-to-pdf")
async def jpg_to_pdf(files: list[UploadFile] = File(...), background_tasks: BackgroundTasks = None):
    image_paths = []
    pdf_path = UPLOAD_DIR / f"jpg-to-pdf_{int(time.time())}.pdf"
    try:
        for file in files:
            img_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
            with open(img_path, "wb") as buffer:
                shutil.copyfileobj(file.file, buffer)
            image_paths.append(str(img_path))
        images = [Image.open(p).convert("RGB") for p in image_paths]
        images[0].save(pdf_path, save_all=True, append_images=images[1:])
        def cleanup():
            for p in image_paths:
                if os.path.exists(p): os.remove(p)
            if os.path.exists(pdf_path): os.remove(pdf_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(pdf_path, filename="jpg-to-pdf.pdf", media_type="application/pdf")
    except Exception as e:
        for p in image_paths:
            if os.path.exists(p): os.remove(p)
        if os.path.exists(pdf_path): os.remove(pdf_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})

@app.post("/api/png-to-pdf")
async def png_to_pdf(files: list[UploadFile] = File(...), background_tasks: BackgroundTasks = None):
    image_paths = []
    pdf_path = UPLOAD_DIR / f"png-to-pdf_{int(time.time())}.pdf"
    try:
        for file in files:
            img_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
            with open(img_path, "wb") as buffer:
                shutil.copyfileobj(file.file, buffer)
            image_paths.append(str(img_path))
        images = [Image.open(p).convert("RGB") for p in image_paths]
        images[0].save(pdf_path, save_all=True, append_images=images[1:])
        def cleanup():
            for p in image_paths:
                if os.path.exists(p): os.remove(p)
            if os.path.exists(pdf_path): os.remove(pdf_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(pdf_path, filename="png-to-pdf.pdf", media_type="application/pdf")
    except Exception as e:
        for p in image_paths:
            if os.path.exists(p): os.remove(p)
        if os.path.exists(pdf_path): os.remove(pdf_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})

@app.post("/api/image-convert")
async def image_convert(file: UploadFile = File(...), format: str = Form(...), background_tasks: BackgroundTasks = None):
    input_path = UPLOAD_DIR / f"{int(time.time())}_{file.filename}"
    output_path = UPLOAD_DIR / f"converted_{int(time.time())}.{format}"
    try:
        with open(input_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        img = Image.open(input_path)
        if format == 'jpg':
            img = img.convert('RGB')
            img.save(output_path, 'JPEG', quality=90)
        elif format == 'png':
            img.save(output_path, 'PNG')
        elif format == 'webp':
            img.save(output_path, 'WEBP')
        else:
            raise ValueError('Unsupported format')
        def cleanup():
            if os.path.exists(input_path): os.remove(input_path)
            if os.path.exists(output_path): os.remove(output_path)
        if background_tasks is not None:
            background_tasks.add_task(cleanup)
        return FileResponse(output_path, filename=f"converted.{format}")
    except Exception as e:
        if os.path.exists(input_path): os.remove(input_path)
        if os.path.exists(output_path): os.remove(output_path)
        return JSONResponse(status_code=500, content={"error": "Conversion failed", "details": str(e)})
